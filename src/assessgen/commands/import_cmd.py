"""import 命令实现 — xlsx → JSON（render 的逆操作）。"""
from pathlib import Path
from typing import Any, Optional

import typer
from loguru import logger

from assessgen.presets.loader import load_preset
from assessgen.utils.io import write_json


def import_cmd(
    file_path: Path = typer.Argument(help="授课计划 xlsx 文件路径"),
    preset: str = typer.Option(..., "--preset", "-p", help="预设名称（提供列映射）"),
    output: Optional[Path] = typer.Option(None, "--output", "-o", help="输出 JSON 路径"),
):
    """将 xlsx 转回 schedule.json（基于 preset 的 style_rules.columns 反向映射）。"""
    if not file_path.exists():
        typer.echo(f"错误: 文件不存在: {file_path}", err=True)
        raise typer.Exit(1)

    try:
        config = load_preset(preset)
    except FileNotFoundError as e:
        typer.echo(f"错误: {e}", err=True)
        raise typer.Exit(1)

    rules = config.style_rules or {}
    columns = rules.get("columns", [])
    data_path = rules.get("data_path")
    if not columns or not data_path:
        typer.echo("错误: preset 缺少 style_rules.columns 或 data_path，无法 import", err=True)
        raise typer.Exit(1)

    entries, shared_fields = _read_xlsx(file_path, columns)

    content: dict[str, Any] = {}
    # 写入 fill_from 字段（按 JSON path）
    for path, value in shared_fields.items():
        _set_value_by_path(content, path, value)
    # 写入 entries
    _set_value_by_path(content, data_path, entries)

    if output is None:
        output = file_path.with_suffix(".json")
    write_json(content, output)
    logger.success(f"JSON 导入成功: {output}（{len(entries)} 条记录）")


def _read_xlsx(
    file_path: Path,
    columns: list[dict],
) -> tuple[list[dict], dict[str, Any]]:
    """读取 xlsx，返回 (entries 列表, fill_from 字段映射)。"""
    from openpyxl import load_workbook

    wb = load_workbook(file_path, data_only=True)
    ws = wb.active

    entries: list[dict] = []
    shared_fields: dict[str, Any] = {}

    # 数据行：从第 2 行开始（第 1 行是表头）
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or all(cell is None or cell == "" for cell in row):
            continue  # 跳过空行

        entry: dict = {}
        for col_idx, col in enumerate(columns):
            if col_idx >= len(row):
                continue
            value = row[col_idx]
            if col.get("auto_increment"):
                continue  # 序号列不入 JSON
            fill_from = col.get("fill_from")
            if fill_from:
                # 共享字段：只记录一次（以首个非空行的值为准）
                if fill_from not in shared_fields and value not in (None, ""):
                    shared_fields[fill_from] = value
                continue
            key = col.get("key")
            if key:
                entry[key] = value if value is not None else ""
        if entry:
            entries.append(entry)

    return entries, shared_fields


def _set_value_by_path(target: dict, path: str, value: Any) -> None:
    """按 `.` 分隔的 JSON path 在 target 中创建嵌套结构并写入 value。"""
    if not path:
        return
    parts = path.split(".")
    current = target
    for key in parts[:-1]:
        if key not in current or not isinstance(current[key], dict):
            current[key] = {}
        current = current[key]
    current[parts[-1]] = value
